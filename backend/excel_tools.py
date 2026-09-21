"""工作区内的 Excel/CSV 专用工具，避免模型用 Bash 试错读表。"""

from __future__ import annotations

import csv
import json
from datetime import date, datetime
from pathlib import Path
from typing import Any

ALLOWED_TABLE_EXT = {".xlsx", ".csv", ".tsv"}


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _parse_json_list(raw: Any, *, name: str) -> list[Any]:
    if raw is None or raw == "":
        return []
    if isinstance(raw, list):
        return raw
    if isinstance(raw, str):
        text = raw.strip()
        if not text:
            return []
        data = json.loads(text)
        if not isinstance(data, list):
            raise ValueError(f"{name} 必须是 JSON 数组")
        return data
    raise ValueError(f"{name} 格式无效")


def _require_table(path: Path) -> None:
    if not path.exists() or not path.is_file():
        raise FileNotFoundError(f"文件不存在: {path.name}")
    ext = path.suffix.lower()
    if ext == ".xls":
        raise ValueError("不支持旧版 .xls，请另存为 .xlsx 或 .csv 后重试")
    if ext not in ALLOWED_TABLE_EXT:
        raise ValueError(f"仅支持 {', '.join(sorted(ALLOWED_TABLE_EXT))}，收到: {ext or '(无扩展名)'}")


def _csv_dialect(path: Path) -> dict[str, Any]:
    delim = "\t" if path.suffix.lower() == ".tsv" else ","
    return {"delimiter": delim}


def excel_info(path: Path) -> str:
    """工作表名、行列数、表头。"""
    _require_table(path)
    ext = path.suffix.lower()
    if ext in {".csv", ".tsv"}:
        with path.open("r", encoding="utf-8-sig", newline="") as f:
            reader = csv.reader(f, **_csv_dialect(path))
            rows = list(reader)
        if not rows:
            return f"文件: {path.name}\n工作表: Sheet1\n行数: 0\n列数: 0\n表头: (空)"
        headers = [_cell_text(c) for c in rows[0]]
        return (
            f"文件: {path.name}\n"
            f"工作表: Sheet1（CSV 视为单表）\n"
            f"行数: {len(rows)}（含表头）\n"
            f"列数: {len(headers)}\n"
            f"表头: {', '.join(headers) if headers else '(空)'}"
        )

    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("需要安装 openpyxl 才能读取 Excel") from exc

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        lines = [f"文件: {path.name}", f"工作表数: {len(wb.sheetnames)}", "---"]
        for name in wb.sheetnames:
            ws = wb[name]
            rows = ws.max_row or 0
            cols = ws.max_column or 0
            headers: list[str] = []
            for cell in next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ()):
                headers.append(_cell_text(cell))
            lines.append(
                f"工作表: {name}\n"
                f"  行数: {rows}（含表头）\n"
                f"  列数: {cols}\n"
                f"  表头: {', '.join(headers) if headers else '(空)'}"
            )
        return "\n".join(lines)
    finally:
        wb.close()


def excel_preview(path: Path, sheet: str = "", limit: int = 30) -> str:
    """预览前 N 行（含表头）。"""
    _require_table(path)
    limit = max(1, min(int(limit or 30), 200))
    ext = path.suffix.lower()

    if ext in {".csv", ".tsv"}:
        with path.open("r", encoding="utf-8-sig", newline="") as f:
            reader = csv.reader(f, **_csv_dialect(path))
            rows = []
            for i, row in enumerate(reader):
                rows.append(row)
                if i + 1 >= limit:
                    break
        if not rows:
            return f"{path.name} / Sheet1\n(空表)"
        return _format_preview(path.name, "Sheet1", rows, limit)

    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("需要安装 openpyxl 才能读取 Excel") from exc

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        name = (sheet or "").strip() or wb.sheetnames[0]
        if name not in wb.sheetnames:
            return f"工作表不存在: {name}。可用: {', '.join(wb.sheetnames)}"
        ws = wb[name]
        rows: list[list[str]] = []
        for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
            rows.append([_cell_text(c) for c in row])
            if i >= limit:
                break
        if not rows:
            return f"{path.name} / {name}\n(空表)"
        return _format_preview(path.name, name, rows, limit)
    finally:
        wb.close()


def _format_preview(filename: str, sheet: str, rows: list[list[Any]], limit: int) -> str:
    lines = [
        f"文件: {filename}",
        f"工作表: {sheet}",
        f"预览前 {min(len(rows), limit)} 行（含表头）",
        "---",
    ]
    for row in rows:
        cells = [_cell_text(c).replace("\t", " ").replace("\n", " ") for c in row]
        lines.append("\t".join(cells))
    return "\n".join(lines)


def excel_write_sheet(
    path: Path,
    *,
    sheet: str = "Sheet1",
    headers: Any = None,
    rows: Any = None,
) -> str:
    """按表头+行数据写入/覆盖 .xlsx 或 .csv。"""
    ext = path.suffix.lower()
    if ext == ".xls":
        raise ValueError("不支持写入旧版 .xls，请使用 .xlsx 或 .csv")
    if ext not in ALLOWED_TABLE_EXT:
        if not ext:
            path = path.with_suffix(".xlsx")
            ext = ".xlsx"
        else:
            raise ValueError(f"仅支持写入 {', '.join(sorted(ALLOWED_TABLE_EXT))}")

    header_list = [_cell_text(h) for h in _parse_json_list(headers, name="headers")]
    raw_rows = _parse_json_list(rows, name="rows")
    if len(raw_rows) > 50_000:
        raise ValueError("行数过多，上限 50000")
    if not header_list and raw_rows:
        first = raw_rows[0]
        if isinstance(first, dict):
            header_list = [str(k) for k in first.keys()]
        elif isinstance(first, list):
            header_list = [f"列{i+1}" for i in range(len(first))]

    normalized: list[list[str]] = []
    for item in raw_rows:
        if isinstance(item, dict):
            normalized.append([_cell_text(item.get(h, "")) for h in header_list])
        elif isinstance(item, list):
            cells = [_cell_text(c) for c in item]
            if header_list and len(cells) < len(header_list):
                cells.extend([""] * (len(header_list) - len(cells)))
            normalized.append(cells)
        else:
            normalized.append([_cell_text(item)])

    path.parent.mkdir(parents=True, exist_ok=True)
    sheet_name = ((sheet or "Sheet1").strip() or "Sheet1")[:31]

    if ext in {".csv", ".tsv"}:
        with path.open("w", encoding="utf-8-sig", newline="") as f:
            writer = csv.writer(f, **_csv_dialect(path))
            if header_list:
                writer.writerow(header_list)
            writer.writerows(normalized)
        return (
            f"已写入 {path.name}\n"
            f"工作表: {sheet_name}（CSV 单表）\n"
            f"列: {len(header_list)}  数据行: {len(normalized)}\n"
            f"路径: {path}"
        )

    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise RuntimeError("需要安装 openpyxl 才能写入 Excel") from exc

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    if header_list:
        ws.append(header_list)
    for row in normalized:
        ws.append(row)
    wb.save(path)
    return (
        f"已写入 {path.name}\n"
        f"工作表: {sheet_name}\n"
        f"列: {len(header_list)}  数据行: {len(normalized)}\n"
        f"路径: {path}"
    )
