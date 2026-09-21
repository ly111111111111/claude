"""从 Excel/CSV 按「部门受案号」建立案件字段索引。带 mtime 缓存。"""
from __future__ import annotations

import csv
import re
from pathlib import Path
from typing import Any

from backend.case_files import extract_case_digit

try:
    from openpyxl import load_workbook

    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

_excel_cache: dict[str, Any] = {
    "path": None,
    "mtime": None,
    "key_col": None,
    "by_digit": {},
    "columns": [],
}


def _cell_to_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else str(value).strip()
    if isinstance(value, int):
        return str(value)
    return str(value).strip()


def normalize_case_key(value: Any) -> str:
    text = _cell_to_str(value)
    if not text:
        return ""
    text = re.sub(r"\s+", "", text)
    text = text.removesuffix("号")
    digit = extract_case_digit(text)
    if digit:
        return digit
    if re.fullmatch(r"\d{8,16}", text):
        return text
    return ""


def _is_csv(path: Path) -> bool:
    return path.suffix.lower() in {".csv", ".tsv"}


def _read_rows(path: Path) -> list[list[str]]:
    if _is_csv(path):
        with path.open(encoding="utf-8-sig") as f:
            reader = csv.reader(f)
            return [[_cell_to_str(c) for c in row] for row in reader]
    if not HAS_OPENPYXL:
        raise RuntimeError("需要安装 openpyxl 才能读取 Excel 文件")
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        return [[_cell_to_str(c) for c in row] for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def _build_index(rows: list[list[str]], key_col: str) -> tuple[list[str], dict[str, dict[str, str]]]:
    if not rows:
        return [], {}
    header = rows[0]
    columns = [h or f"列{i+1}" for i, h in enumerate(header)]
    key_idx = None
    for i, name in enumerate(columns):
        if name == key_col or key_col in name:
            key_idx = i
            break
    if key_idx is None:
        raise ValueError(f"未找到列「{key_col}」，实际列: {', '.join(columns)}")
    by_digit: dict[str, dict[str, str]] = {}
    for row in rows[1:]:
        if not row:
            continue
        key = normalize_case_key(row[key_idx] if key_idx < len(row) else "")
        if not key:
            continue
        record = {col: row[i] if i < len(row) else "" for i, col in enumerate(columns)}
        by_digit.setdefault(key, record)
    return columns, by_digit


def load_case_excel_index(path: Path, key_col: str = "部门受案号") -> dict[str, Any]:
    path = path.expanduser().resolve()
    if not path.exists() or not path.is_file():
        raise FileNotFoundError(f"案件文件不存在: {path}")
    try:
        mtime = path.stat().st_mtime
    except OSError as exc:
        raise FileNotFoundError(f"案件文件不可读: {path}") from exc
    if (
        _excel_cache["path"] == str(path)
        and _excel_cache["mtime"] == mtime
        and _excel_cache["key_col"] == key_col
    ):
        return {
            "path": str(path),
            "columns": _excel_cache["columns"],
            "by_digit": _excel_cache["by_digit"],
            "count": len(_excel_cache["by_digit"]),
        }
    rows = _read_rows(path)
    columns, by_digit = _build_index(rows, key_col)
    _excel_cache.update(path=str(path), mtime=mtime, key_col=key_col, columns=columns, by_digit=by_digit)
    return {"path": str(path), "columns": columns, "by_digit": by_digit, "count": len(by_digit)}


def lookup_case_excel(path: Path | None, query: str, key_col: str = "部门受案号") -> dict[str, Any] | None:
    if path is None:
        return None
    index = load_case_excel_index(path, key_col)
    digit = normalize_case_key(query) or extract_case_digit(query)
    if not digit:
        return None
    row = index["by_digit"].get(digit)
    if not row:
        return None
    fields = [{"key": k, "value": row.get(k, "")} for k in index["columns"]]
    return {"digit": digit, "fields": fields, "row": row, "source": index["path"]}


def list_excel_digits(path: Path | None, key_col: str = "部门受案号") -> list[str]:
    if path is None:
        return []
    index = load_case_excel_index(path, key_col)
    return sorted(index["by_digit"].keys())
